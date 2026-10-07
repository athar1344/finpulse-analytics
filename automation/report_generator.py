"""Generate Excel, PDF, and CSV reports from the processed dataset."""
import os
import sys

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

import config
from utils.helper import get_logger

logger = get_logger("report_generator")


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def build_summary(df: pd.DataFrame) -> pd.DataFrame:
    actuals = df[df["Scenario"] == "Actuals"]
    summary = (
        actuals.groupby(["business_unit", "Year"])["annual_total"]
        .sum()
        .reset_index()
        .rename(columns={"annual_total": "total_actuals"})
    )
    return summary


def export_csv(df: pd.DataFrame) -> str:
    path = os.path.join(config.REPORTS_DIR, f"financial_report_{_timestamp()}.csv")
    df.to_csv(path, index=False)
    logger.info("CSV report generated: %s", path)
    return path


def export_excel(df: pd.DataFrame) -> str:
    path = os.path.join(config.REPORTS_DIR, f"financial_report_{_timestamp()}.xlsx")
    summary = build_summary(df)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Raw Data", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)

        workbook = writer.book
        header_fill = PatternFill(start_color="0068C9", end_color="0068C9", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for sheet_name in ("Raw Data", "Summary"):
            sheet = workbook[sheet_name]
            for col_idx, cell in enumerate(sheet[1], start=1):
                cell.fill = header_fill
                cell.font = header_font
                col_letter = get_column_letter(col_idx)
                sheet.column_dimensions[col_letter].width = 16

    logger.info("Excel report generated: %s", path)
    return path


def export_pdf(df: pd.DataFrame) -> str:
    path = os.path.join(config.REPORTS_DIR, f"financial_report_{_timestamp()}.pdf")
    summary = build_summary(df)

    doc = SimpleDocTemplate(path, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Financial Summary Report", styles["Title"]),
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Spacer(1, 16),
    ]

    table_data = [list(summary.columns)] + summary.round(2).astype(str).values.tolist()
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0068C9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)

    logger.info("PDF report generated: %s", path)
    return path


def generate_all_reports(df: pd.DataFrame) -> dict:
    return {
        "csv": export_csv(df),
        "excel": export_excel(df),
        "pdf": export_pdf(df),
    }


if __name__ == "__main__":
    data = pd.read_csv(config.PROCESSED_FILE_PATH)
    print(generate_all_reports(data))
